import pptx
from pptx import Presentation

template_path = r'C:\Users\kiran akash\Downloads\SIH2026-IDEA-Presentation-Format.pptx'
prs = Presentation(template_path)

print(f"Total slides: {len(prs.slides)}")
print(f"Slide dimensions: {prs.slide_width / 914400:.2f} x {prs.slide_height / 914400:.2f} inches")

for idx, slide in enumerate(prs.slides):
    print(f"\n=== SLIDE {idx+1} ===")
    for shape in slide.shapes:
        left = shape.left / 914400
        top = shape.top / 914400
        width = shape.width / 914400
        height = shape.height / 914400
        print(f"  Shape [{shape.name}] (type={type(shape).__name__}) at ({left:.2f}, {top:.2f}) size ({width:.2f} x {height:.2f})")
        if shape.has_text_frame:
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                fonts = [r.font.name for r in p.runs if r.font.name]
                sizes = [r.font.size.pt for r in p.runs if r.font.size]
                bold = [r.font.bold for r in p.runs]
                print(f"    P{p_idx}: {repr(p.text)} | font={fonts} size={sizes} bold={bold}")
