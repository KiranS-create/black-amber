"""
SIH26237 — Official SIH 2026 PPTX Template Integrity Test Suite
================================================================

Validates:
1. presentation/SIH26237_Final.pptx exists and parses with python-pptx.
2. Exactly 6 slides (SIH maximum limit).
3. 16:9 widescreen aspect ratio.
4. Official SIH branding elements, footer bar, and logo pictures are preserved.
5. All 6 slides contain verified SIH26237 content with zero remaining template placeholder text.
6. Original template file was preserved untouched.
"""

from pathlib import Path
import pytest
import pptx
from pptx import Presentation

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
FINAL_PPTX_PATH = PROJECT_ROOT / "presentation" / "SIH26237_Final.pptx"
TEMPLATE_PATH = Path(r"C:\Users\kiran akash\Downloads\SIH2026-IDEA-Presentation-Format.pptx")

def test_final_pptx_exists_and_parses():
    assert FINAL_PPTX_PATH.exists(), f"Final PPTX does not exist at {FINAL_PPTX_PATH}"
    assert FINAL_PPTX_PATH.stat().st_size > 50000, "PPTX file size is suspiciously small"
    prs = Presentation(str(FINAL_PPTX_PATH))
    assert prs is not None

def test_exact_six_slide_count():
    prs = Presentation(str(FINAL_PPTX_PATH))
    assert len(prs.slides) == 6, f"Expected exactly 6 slides (SIH limit), found {len(prs.slides)}"

def test_widescreen_aspect_ratio():
    prs = Presentation(str(FINAL_PPTX_PATH))
    ratio = prs.slide_width / prs.slide_height
    # 16:9 ratio is 1.777777...
    assert 1.75 <= ratio <= 1.80, f"Aspect ratio {ratio:.3f} is not 16:9 widescreen"

def test_required_sih_branding_elements():
    prs = Presentation(str(FINAL_PPTX_PATH))
    for idx, slide in enumerate(prs.slides):
        # Every slide must have shapes
        assert len(slide.shapes) >= 4, f"Slide {idx+1} has too few shapes"
        
        # Check for pictures (SIH Logo) on each slide
        pictures = [s for s in slide.shapes if type(s).__name__ == "Picture"]
        assert len(pictures) >= 1, f"Slide {idx+1} is missing SIH branding logo image"
        
        # Check for footer on content slides 2..6
        if idx > 0:
            footer_texts = [s.text for s in slide.shapes if s.has_text_frame and "@SIH Idea submission" in s.text]
            assert len(footer_texts) >= 1, f"Slide {idx+1} missing SIH footer placeholder"

def test_slide_titles_and_content():
    prs = Presentation(str(FINAL_PPTX_PATH))
    
    # Slide 1 (Title Page)
    slide1_text = " ".join([s.text for s in prs.slides[0].shapes if s.has_text_frame])
    assert "SIH26237" in slide1_text
    assert "Post-Quantum Document Leak Attribution" in slide1_text
    assert "Ad Astra" in slide1_text or "iTantra" in slide1_text
    
    # Slide 2 (Proposed Solution)
    slide2_text = " ".join([s.text for s in prs.slides[1].shapes if s.has_text_frame])
    assert "PROPOSED SOLUTION" in slide2_text
    assert "ML-KEM-768" in slide2_text
    assert "Tardos" in slide2_text
    
    # Slide 3 (Technical Approach)
    slide3_text = " ".join([s.text for s in prs.slides[2].shapes if s.has_text_frame])
    assert "TECHNICAL APPROACH" in slide3_text
    assert "ML-DSA-65" in slide3_text
    assert "Evidence Dependency Graph" in slide3_text
    
    # Slide 4 (Feasibility and Viability)
    slide4_text = " ".join([s.text for s in prs.slides[3].shapes if s.has_text_frame])
    assert "FEASIBILITY AND VIABILITY" in slide4_text
    assert "100% Offline" in slide4_text or "Offline" in slide4_text
    assert "simulated" in slide4_text.lower()
    
    # Slide 5 (Impact and Benefits)
    slide5_text = " ".join([s.text for s in prs.slides[4].shapes if s.has_text_frame])
    assert "IMPACT AND BENEFITS" in slide5_text
    assert "11.61 ms" in slide5_text
    assert "74.0%" in slide5_text or "74%" in slide5_text
    
    # Slide 6 (Research and References)
    slide6_text = " ".join([s.text for s in prs.slides[5].shapes if s.has_text_frame])
    assert "RESEARCH AND REFERENCES" in slide6_text
    assert "NIST FIPS 203" in slide6_text
    assert "Fail-Closed" in slide6_text or "fail-closed" in slide6_text

def test_no_raw_template_placeholders():
    prs = Presentation(str(FINAL_PPTX_PATH))
    raw_placeholders = [
        "Detailed explanation of the proposed solution",
        "How it addresses the problem",
        "Innovation and uniqueness of the solution",
        "Technologies to be used (e.g. programming languages",
        "Methodology and process for implementation",
        "Analysis of the feasibility of the idea",
        "Potential challenges and risks",
        "Details / Links of the reference and research work"
    ]
    for idx, slide in enumerate(prs.slides):
        slide_text = " ".join([s.text for s in slide.shapes if s.has_text_frame])
        for ph in raw_placeholders:
            assert ph not in slide_text, f"Slide {idx+1} contains raw template placeholder text: '{ph}'"

def test_original_template_unmodified():
    assert TEMPLATE_PATH.exists(), "Original template was deleted or moved!"
    prs = Presentation(str(TEMPLATE_PATH))
    # Original raw template has 7 slides (including instruction slide)
    assert len(prs.slides) == 7, "Original template was altered!"
